"""Build Project_Report.docx and Project_Report.pdf from Project_Report.md.

    python docs/build_report.py

Requires: python-docx, reportlab   (pip install python-docx reportlab)
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (Image, PageBreak, Paragraph, Preformatted,
                                SimpleDocTemplate, Spacer, Table, TableStyle)

HERE = Path(__file__).resolve().parent
MD = HERE / "Project_Report.md"
TITLE = "Iron-Watch: Metal-Detecting Access Gate with Face-Recognition Attendance Logging"

# ------------------------------------------------------------------ parse md
def parse(md_text):
    """Yield (kind, payload) blocks: h1/h2/h3, p, ul, table, code, img, rule."""
    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("```"):
            buf = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                buf.append(lines[i]); i += 1
            i += 1
            yield "code", "\n".join(buf); continue
        if ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            yield "table", rows; continue
        m = re.match(r"!\[(.*?)\]\((.*?)\)", ln.strip())
        if m:
            yield "img", (m.group(1), HERE / m.group(2)); i += 1; continue
        if ln.startswith("#"):
            lvl = len(ln) - len(ln.lstrip("#"))
            yield f"h{min(lvl,3)}", ln.lstrip("#").strip(); i += 1; continue
        if ln.strip() == "---":
            yield "rule", None; i += 1; continue
        if re.match(r"^\s*([-*]|\d+\.)\s", ln):
            items = []
            while i < len(lines) and re.match(r"^\s*([-*]|\d+\.)\s", lines[i]):
                items.append(re.sub(r"^\s*([-*]|\d+\.)\s", "", lines[i])); i += 1
            yield "ul", items; continue
        if ln.strip() == "":
            i += 1; continue
        buf = [ln]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#|\||```|!\[|---|\s*[-*]\s|\s*\d+\.\s)", lines[i]):
            buf.append(lines[i]); i += 1
        yield "p", " ".join(buf)

# ------------------------------------------------------------------ inline md
def strip_inline(t):
    t = re.sub(r"<sub>(.*?)</sub>", r"\1", t)
    t = re.sub(r"\*\*(.*?)\*\*", r"\1", t)
    t = re.sub(r"(?<!\*)\*(?!\*)(.*?)\*", r"\1", t)
    t = re.sub(r"`(.*?)`", r"\1", t)
    return t

def rl_inline(t):
    t = t.replace("&", "&amp;").replace("<sub>", "\x01").replace("</sub>", "\x02")
    t = t.replace("<", "&lt;").replace(">", "&gt;")
    t = t.replace("\x01", "<sub>").replace("\x02", "</sub>")
    codes = []
    def keep(m):
        codes.append(m.group(1)); return f"\x03{len(codes)-1}\x03"
    t = re.sub(r"`(.*?)`", keep, t)
    t = re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", t)
    t = re.sub(r"(?<!\*)\*(?!\*)(.*?)\*", r"<i>\1</i>", t)
    t = re.sub(r"\x03(\d+)\x03", lambda m: f"<font face='Courier'>{codes[int(m.group(1))]}</font>", t)
    t = re.sub(r"\[(.*?)\]\((.*?)\)", r"\1", t)
    return t

def docx_runs(par, t):
    t = re.sub(r"\[(.*?)\]\((.*?)\)", r"\1", t)
    for tok in re.split(r"(\*\*.*?\*\*|`.*?`|<sub>.*?</sub>|(?<!\*)\*(?!\*).*?\*)", t):
        if not tok:
            continue
        if tok.startswith("**"):
            par.add_run(strip_inline(tok)).bold = True
        elif tok.startswith("`"):
            r = par.add_run(tok[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9.5)
        elif tok.startswith("<sub>"):
            par.add_run(tok[5:-6]).font.subscript = True
        elif tok.startswith("*"):
            par.add_run(tok[1:-1]).italic = True
        else:
            par.add_run(tok)

# ------------------------------------------------------------------ DOCX
def build_docx(blocks, out):
    d = Document()
    st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    for s in d.sections:
        s.left_margin = s.right_margin = Inches(1); s.top_margin = s.bottom_margin = Inches(1)
    first_h1 = True
    for kind, pl in blocks:
        if kind == "h1":
            if first_h1:   # cover page
                for _ in range(6): d.add_paragraph()
                p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(pl); r.bold = True; r.font.size = Pt(22)
                first_h1 = False
            else:
                d.add_page_break(); d.add_heading(pl, 1)
        elif kind == "h2":
            if pl.strip() in ("Abstract", "1 Introduction"):
                d.add_page_break()
            d.add_heading(strip_inline(pl), 1)
        elif kind == "h3":
            d.add_heading(strip_inline(pl), 2)
        elif kind == "p":
            if pl.startswith("**") and pl.endswith("**") and "—" in pl:
                p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(strip_inline(pl)); r.bold = True; r.font.size = Pt(13)
            else:
                docx_runs(d.add_paragraph(), pl)
        elif kind == "ul":
            for it in pl:
                docx_runs(d.add_paragraph(style="List Bullet"), it)
        elif kind == "code":
            p = d.add_paragraph(); r = p.add_run(pl); r.font.name = "Consolas"; r.font.size = Pt(8.5)
        elif kind == "table":
            rows = pl; ncol = max(len(r) for r in rows)
            t = d.add_table(rows=len(rows), cols=ncol); t.style = "Light Grid Accent 1"
            for i, row in enumerate(rows):
                for j in range(ncol):
                    cell = t.cell(i, j); cell.text = ""
                    docx_runs(cell.paragraphs[0], row[j] if j < len(row) else "")
                    for r in cell.paragraphs[0].runs:
                        r.font.size = Pt(9.5)
                        if i == 0: r.bold = True
            d.add_paragraph()
        elif kind == "img":
            cap, path = pl
            if path.exists():
                d.add_picture(str(path), width=Inches(6))
                d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif kind == "rule":
            pass
    d.save(out)

# ------------------------------------------------------------------ PDF
def build_pdf(blocks, out):
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontSize=10.5, leading=14, spaceAfter=6)
    cell = ParagraphStyle("cell", parent=body, fontSize=8.5, leading=10.5, spaceAfter=0)
    cap = ParagraphStyle("cap", parent=body, fontSize=9, alignment=TA_CENTER, textColor=colors.grey)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontSize=16, spaceBefore=14, spaceAfter=8, textColor=colors.HexColor("#1f2937"))
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontSize=12.5, spaceBefore=10, spaceAfter=5)
    h3 = ParagraphStyle("h3", parent=ss["Heading3"], fontSize=11, spaceBefore=8, spaceAfter=4)
    cover = ParagraphStyle("cover", parent=body, fontSize=20, leading=26, alignment=TA_CENTER, spaceAfter=18)
    sub = ParagraphStyle("sub", parent=body, fontSize=12, alignment=TA_CENTER, spaceAfter=10)
    code = ParagraphStyle("code", parent=ss["Code"], fontSize=7.8, leading=9.5, backColor=colors.HexColor("#f3f4f6"), borderPadding=4, spaceAfter=8)

    story = []
    first_h1 = True
    W = A4[0] - 4 * cm
    for kind, pl in blocks:
        if kind == "h1":
            if first_h1:
                story += [Spacer(1, 5 * cm), Paragraph(pl, cover)]
                first_h1 = False
            else:
                story += [PageBreak(), Paragraph(rl_inline(pl), h1)]
        elif kind == "h2":
            if pl.strip() in ("Abstract", "1 Introduction"):
                story.append(PageBreak())
            story.append(Paragraph(rl_inline(pl), h1))
        elif kind == "h3":
            story.append(Paragraph(rl_inline(pl), h2))
        elif kind == "p":
            style = sub if (pl.startswith("**") and pl.endswith("**") and "—" in pl) else \
                    cap if (pl.startswith("*Figure") or pl.startswith("*(")) else body
            story.append(Paragraph(rl_inline(pl), style))
        elif kind == "ul":
            for it in pl:
                story.append(Paragraph("• " + rl_inline(it), body))
        elif kind == "code":
            story.append(Preformatted(pl, code))
        elif kind == "table":
            rows = pl; ncol = max(len(r) for r in rows)
            data = [[Paragraph(rl_inline(r[j]) if j < len(r) else "", cell) for j in range(ncol)] for r in rows]
            t = Table(data, colWidths=[W / ncol] * ncol, repeatRows=1)
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9ca3af")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f4f6")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ]))
            story += [t, Spacer(1, 8)]
        elif kind == "img":
            _, path = pl
            if path.exists():
                img = Image(str(path)); ratio = img.imageHeight / img.imageWidth
                w = min(W, 15 * cm); img.drawWidth = w; img.drawHeight = w * ratio
                story.append(img)
        elif kind == "rule":
            pass

    def footer(canvas, doc):
        canvas.saveState(); canvas.setFont("Helvetica", 8); canvas.setFillColor(colors.grey)
        canvas.drawString(2 * cm, 1.2 * cm, "Iron-Watch — AASTU Electromechanical Engineering internship project")
        canvas.drawRightString(A4[0] - 2 * cm, 1.2 * cm, f"Page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(str(out), pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm,
                            topMargin=2 * cm, bottomMargin=2 * cm, title=TITLE, author="Iron-Watch")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


if __name__ == "__main__":
    blocks = list(parse(MD.read_text(encoding="utf-8")))
    build_docx(blocks, HERE / "Project_Report.docx")
    build_pdf(blocks, HERE / "Project_Report.pdf")
    print("Wrote Project_Report.docx and Project_Report.pdf")
